"""Generate review fixtures from HD05, independent of the checker's rule tables."""
import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Mm, Pt


def make_party_sample(output, *, letter=False, errors=False):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin, section.right_margin = Mm(30), Mm(15)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(14)
    normal.font.bold = normal.font.italic = False

    def line(parent, text, size=14, bold=False, italic=False, underline=False):
        paragraph = parent.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.size, run.bold, run.italic, run.underline = Pt(size), bold, italic, underline
        return paragraph

    header = document.add_table(rows=1, cols=2)
    left, right = header.rows[0].cells
    line(left, "ĐẢNG BỘ THÀNH PHỐ HÀ NỘI")
    line(left, "ĐẢNG UỶ PHƯỜNG BA ĐÌNH", bold=True)
    line(left, "*")
    code = "CV" if letter else "KH"
    line(left, f"Số: 5-{code}/ĐU" if errors else f"Số 05-{code}/ĐU")
    if letter:
        line(left, "V/v triển khai nhiệm vụ công tác" if errors else "về việc triển khai nhiệm vụ công tác", 13 if errors else 12, italic=not errors)
        line(left, "năm 2026", 12, italic=True)
    line(right, "ĐẢNG CỘNG SẢN VIỆT NAM", 15, bold=True)
    line(right, "Hà Nội, ngày 03 tháng 02 năm 2026", 13 if errors else 14, italic=True)
    if not letter:
        line(document, "KẾ HOẠCH", 15, bold=True)
        line(document, "triển khai nhiệm vụ công tác năm 2026", 14, bold=True)
        line(document, "-----")
    else:
        greeting = document.add_paragraph()
        greeting.alignment = WD_ALIGN_PARAGRAPH.CENTER
        greeting.add_run("Kính gửi: ").italic = True
        greeting.add_run("Các chi bộ trực thuộc;").italic = False
    for text in (
        "Các chi bộ triển khai nhiệm vụ công tác theo kế hoạch đã được thông qua.",
        "Các đơn vị tổ chức thực hiện và báo cáo kết quả về cơ quan theo quy định.",
    ):
        paragraph = document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Mm(12.7 if errors else 10)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = Pt(20)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    signature = document.add_table(rows=1, cols=2)
    left, right = signature.rows[0].cells
    line(left, "Nơi nhận:", underline=True)
    line(left, "- Các chi bộ trực thuộc;", 11 if errors else 12)
    line(left, "- Lưu Văn phòng Đảng uỷ.", 12)
    line(right, "T/M ĐẢNG UỶ", bold=True)
    line(right, "BÍ THƯ", bold=errors)
    line(right, "Nguyễn Văn An", bold=True)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="data/sample-party.docx")
    parser.add_argument("--letter", action="store_true")
    parser.add_argument("--errors", action="store_true")
    options = parser.parse_args()
    make_party_sample(options.out, letter=options.letter, errors=options.errors)
