"""Verify actual CLI/GUI processes and their Word exports before packaging.

With --cli/--gui this checks the frozen programs, not the imported source.
Reports and sample documents stay in the chosen output directory.
"""
import argparse
import hashlib
import json
import os
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt

from make_party_sample import make_party_sample
from make_sample import make_sample
from validate_model import CLEAN, REQUIRED
from vietspell.docx_io.reader import read_docx


def verify(output, cli=None, gui=None, allow_demo=False):
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    command = [str(Path(cli).resolve())] if cli else [sys.executable, "-m", "vietspell"]
    report = {"passed": False, "created_utc": datetime.now(timezone.utc).isoformat(),
              "platform": sys.platform, "target": "frozen" if cli else "source", "checks": []}

    def run(arguments, expected=0, *, executable=None, environment=None):
        result = subprocess.run((executable or command)+arguments, capture_output=True,
                                encoding="utf-8", errors="replace", timeout=90, env=environment)
        if result.returncode != expected:
            raise ValueError(f"{arguments[0]}: mã {result.returncode}, cần {expected}: {result.stderr[-2000:]} {result.stdout[-2000:]}")
        return result.stdout

    def checked(name):
        report["checks"].append({"name": name, "passed": True})
        print("PASS", name, flush=True)

    def check_file(path, profile):
        target_report = output / f"{path.stem}.json"
        # Either clean or findings is valid here; determine the exit status from
        # the generated JSON, then verify it agrees with the CLI contract.
        result = subprocess.run(command+["check", str(path), "--profile", profile, "--report", str(target_report)],
                                capture_output=True, encoding="utf-8", errors="replace", timeout=90)
        if result.returncode not in (0, 1):
            raise ValueError(f"Soát {path.name}: {result.stderr[-2000:]}")
        payload = json.loads(target_report.read_text(encoding="utf-8"))
        if result.returncode != int(bool(payload["issues"])):
            raise ValueError("Mã thoát check không khớp báo cáo.")
        return payload

    try:
        doctor = json.loads(run(["doctor"]))
        report["model"] = doctor
        quality = doctor["model"].get("quality")
        if quality != "corpus" and not allow_demo:
            raise ValueError("Bản phát hành cần mô hình corpus; mô hình mẫu phải được cho phép riêng.")
        checked("SQLite doctor / model quality")
        if quality == "corpus":
            for index, (text, expected) in enumerate(REQUIRED):
                payload = json.loads(run(["check", "--text", text], expected=1))
                actual = [(i["text"], i["suggestions"][0]["text"]) for i in payload["issues"] if i["suggestions"]]
                if actual != expected:
                    raise ValueError(f"Gợi ý không khớp ca {index}: {actual}")
                checked(f"language regression {index+1}")
            for index, text in enumerate(CLEAN):
                if json.loads(run(["check", "--text", text]))["issues"]:
                    raise ValueError(f"Báo nhầm ca sạch {index}")
                checked(f"clean language {index+1}")
        sample = output / "sample-hc.docx"
        make_sample(sample)
        party_bad = output / "party-letter-errors.docx"
        make_party_sample(party_bad, letter=True, errors=True)
        required_rules = {"PARTY_NUMBER_FORMAT", "PARTY_NO_VV", "ELEMENT_SIZE_DATE",
                          "ELEMENT_SIZE_LETTER_SUBJECT", "ELEMENT_ITALIC_LETTER_SUBJECT",
                          "ELEMENT_BOLD_POSITION", "ELEMENT_SIZE_RECIPIENTS", "FIRST_LINE_INDENT"}
        party_issues = check_file(party_bad, "DANG-HD05")["issues"]
        if not required_rules <= {i["rule_id"] for i in party_issues}:
            raise ValueError("Thiếu phát hiện ở văn bản Đảng cố ý sai.")
        checked("HD05 incorrect letter")
        for letter in (False, True):
            clean = output / ("party-letter-clean.docx" if letter else "party-plan-clean.docx")
            make_party_sample(clean, letter=letter)
            clean_result = check_file(clean, "DANG-HD05")
            if clean_result["issues"]:
                raise ValueError(f"Mẫu HD05 đúng bị báo lỗi: {clean_result['issues']}")
            checked(f"HD05 clean {'letter' if letter else 'plan'}")
        # Exercise the new element ranges through the real CLI too, including
        # frozen Windows builds. Unit tests alone do not test the packaged code.
        for profile in ("HC-ND30", "DANG-HD05"):
            combined = output / f"combined-header-{profile}.docx"
            if profile == "HC-ND30":
                make_sample(combined)
            else:
                make_party_sample(combined)
            document = Document(combined)
            left, right = document.tables[0].rows[0].cells
            number = next(p for p in left.paragraphs if p.text.startswith("Số"))
            date = next(p for p in right.paragraphs if "ngày" in p.text)
            number_text, date_text = number.text, date.text
            date.clear()
            number.clear()
            size = 13 if profile == "HC-ND30" else 14
            number_run = number.add_run(number_text)
            number_run.font.size, number_run.italic, number_run.bold = Pt(size), False, False
            number.add_run("\t")
            date_run = number.add_run(date_text)
            date_run.font.size, date_run.italic, date_run.bold = Pt(size), True, False
            if profile == "HC-ND30":
                national, motto = right.paragraphs[:2]
                national.text = national.text.upper()
                national.runs[0].font.size = Pt(12)
                motto.runs[0].font.size = Pt(13)
            document.save(combined)
            header_findings = [i for i in check_file(combined, profile)["issues"]
                               if i["rule_id"] in {"MISSING_DATE", "MISSING_NUMBER", "ELEMENT_SIZE_CHAIN"}
                               or i["rule_id"].startswith("ELEMENT_") and i["rule_id"].endswith(("_NUMBER", "_DATE"))]
            if header_findings:
                raise ValueError(f"Số/ngày chung đoạn đúng bị báo sai: {profile}: {header_findings}")
            checked(f"combined header independent fonts / {profile}")
            date_run.italic = False
            bad_header = output / f"combined-header-{profile}-wrong-date.docx"
            document.save(bad_header)
            findings = check_file(bad_header, profile)["issues"]
            italic = [i for i in findings if i["rule_id"] in {"ELEMENT_ITALIC_NUMBER", "ELEMENT_ITALIC_DATE"}]
            if [(i["rule_id"], i["text"]) for i in italic] != [("ELEMENT_ITALIC_DATE", date_text)]:
                raise ValueError(f"Phạm vi cảnh báo ngày chưa đúng: {profile}: {italic}")
            checked(f"combined header date error span / {profile}")
        for source, profile in ((sample, "HC-ND30"), (party_bad, "DANG-HD05")):
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            original_text = read_docx(source).text
            check_file(source, profile)
            for mode in ("RED", "TRACK", "PLAIN", "HIGHLIGHT"):
                exported = output / f"{source.stem}-{mode}.docx"
                run(["fix", str(source), "--profile", profile, "-o", str(exported), "--mode", mode])
                result = read_docx(exported)
                if mode == "HIGHLIGHT":
                    valid = result.text == original_text and bool(result.document.element.xpath('.//w:highlight'))
                elif profile == "HC-ND30":
                    valid = "Hôm nay bạn ăn gì? Tôi đang đi chơi." in result.text
                else:
                    valid = "Số 05-CV/ĐU" in result.text and "về việc triển khai" in result.text and "V/v" not in result.text
                if not valid:
                    raise ValueError(f"Sai nội dung xuất {profile}/{mode}")
                if mode == "TRACK" and not result.document.element.xpath('.//w:ins'):
                    raise ValueError("Bản TRACK thiếu revision XML.")
                if mode == "RED" and not result.document.element.xpath('.//w:color[@w:val="FF0000"]'):
                    raise ValueError("Bản RED thiếu màu đỏ.")
                checked(f"Word {profile} / {mode}")
            if hashlib.sha256(source.read_bytes()).hexdigest() != original_hash:
                raise ValueError("Tài liệu gốc bị thay đổi.")
            checked(f"original unchanged / {profile}")
        environment = os.environ.copy()
        environment["QT_QPA_PLATFORM"] = "offscreen"
        gui_command = [str(Path(gui).resolve())] if gui else [sys.executable, str(Path(__file__).resolve().parents[1] / "packaging/gui_entry.py")]
        for source, profile in ((sample, "HC-ND30"), (party_bad, "DANG-HD05")):
            gui_report = output / f"gui-{profile}.json"
            run(["--smoke-test", str(source), "--profile", profile, "--report", str(gui_report)], executable=gui_command, environment=environment)
            if not json.loads(gui_report.read_text(encoding="utf-8"))["passed"]:
                raise ValueError(f"GUI worker lỗi: {profile}")
            checked(f"GUI worker / {profile}")
        report["passed"] = True
    except Exception as error:
        report["error"] = str(error)
    finally:
        (output / "release-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="data/release-check")
    parser.add_argument("--cli")
    parser.add_argument("--gui")
    parser.add_argument("--allow-demo", action="store_true")
    options = parser.parse_args()
    result = verify(options.out, options.cli, options.gui, options.allow_demo)
    if not result["passed"]:
        print(result["error"], file=sys.stderr)
    raise SystemExit(0 if result["passed"] else 1)
