import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt


def make_sample(output):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin, section.right_margin = Mm(30), Mm(15)
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(13)
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, lines in zip(table.rows[0].cells, (
        ["BỘ CÔNG AN", "CÔNG AN TỈNH NINH BÌNH", "Số:         /CAT-ANKT"],
        ["Cộng hòa xã hội chủ nghĩa Việt Nam", "Độc lập – Tự do – Hạnh phúc", "Ninh Bình, ngày         tháng 10 năm 2026"]
    )):
        cell.text = lines[0]
        for line in lines[1:]:
            cell.add_paragraph(line)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Kính gửi: Các đơn vị liên quan")
    p = document.add_paragraph()
    p.add_run("Hôm nay bạn ")
    p.add_run("an").bold = True
    p.add_run(" gì. Tôi đang ")
    p.add_run("di").italic = True
    p.add_run(" chơi. Trong công tác, Nhất là công tác giải phóng mặt bằng cần được quan tâm.")
    p.paragraph_format.first_line_indent = Mm(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    document.add_paragraph("Công trình cống Đôi đã hoàn thành.")
    document.add_paragraph("GIÁM ĐỐC")
    document.add_paragraph("Nguyễn Văn Nam")
    document.add_paragraph("Nơi nhận:")
    document.add_paragraph("- Như trên;")
    document.add_paragraph("- Lưu: VT.")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--out", default="data/sample.docx")
    make_sample(p.parse_args().out)
