import time

import pytest
from docx import Document
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QPalette
from PySide6.QtWidgets import QDialog

from vietspell.core.detector import CheckerConfig
from vietspell.gui.app import MainWindow, NoticeBar, SettingsDialog


def wait_until(qapp, predicate, seconds=10):
    deadline = time.monotonic()+seconds
    while not predicate() and time.monotonic() < deadline:
        qapp.processEvents()
        time.sleep(0.002)
    qapp.processEvents()
    assert predicate(), "Qt action did not finish in time"


def luminance(color):
    channels = [color.redF(), color.greenF(), color.blueF()]
    channels = [v/12.92 if v <= 0.04045 else ((v+0.055)/1.055)**2.4 for v in channels]
    return sum(a*b for a, b in zip(channels, [0.2126, 0.7152, 0.0722]))


def test_light_palette_has_wcag_contrast(qapp):
    palette = qapp.palette()
    for foreground, background in ((QPalette.Text, QPalette.Base), (QPalette.WindowText, QPalette.Window), (QPalette.ButtonText, QPalette.Button)):
        values = sorted((luminance(palette.color(foreground)), luminance(palette.color(background))))
        assert (values[1]+0.05)/(values[0]+0.05) >= 7


@pytest.mark.parametrize("state", ["opened", "running", "clean", "light", "issues", "error"])
def test_six_notice_states(qapp, state):
    bar = NoticeBar()
    bar.show_notice(state, "Thông báo <b>chữ thuần</b>")
    assert bar.state == state
    assert bar.label.textFormat() == Qt.PlainText
    assert bar.label.text() == "Thông báo <b>chữ thuần</b>"
    bar.close()


def test_settings_presets_include_diacritic_threshold(qapp):
    dialog = SettingsDialog(CheckerConfig())
    dialog.preset.setCurrentIndex(0)
    assert dialog.config().min_delta_diacritic == 2.5
    assert dialog.config().min_delta_fuzzy == 9.0
    dialog.close()


def test_missing_model_and_bad_document_use_notice_not_popup(qapp, model_dir):
    w = MainWindow(model_dir / "missing")
    assert w.notice.state == "error"
    w.open_file("nonexistent.docx")
    assert w.notice.state == "error"
    assert not any(isinstance(widget, QDialog) and widget.isVisible() for widget in qapp.topLevelWidgets())
    w.close()


def test_background_check_real_document_and_filter(qapp, model_dir, tmp_path):
    document = Document()
    document.add_paragraph("😀 Hôm nay bạn an gì. Tôi đang di chơi.")
    path = tmp_path / "văn bản.docx"
    document.save(path)
    w = MainWindow(model_dir)
    w.show()
    w.open_file(path)
    assert w.notice.state == "opened"
    w.check()
    assert w.worker is not None
    assert w.notice.state == "running"
    assert not w.open_action.isEnabled()
    wait_until(qapp, lambda: w.worker is None)
    assert w.session.result is not None
    assert w.table.rowCount() == len(w.session.result.issues)
    replacements = [(i.text, i.suggestions[0].text) for i in w.session.result.issues if i.editable]
    assert replacements == [("an", "ăn"), (".", "?"), ("di", "đi")]
    editable_row = next(r for r, i in enumerate(w.session.result.issues) if i.text == "an")
    w.table.setCurrentCell(editable_row, 1)
    qapp.processEvents()
    assert w.text.textCursor().selectedText() == "an"  # UTF-16 mapping after emoji.
    manual_rows = [r for r, i in enumerate(w.session.result.issues) if not i.editable]
    assert manual_rows
    assert all(not w.table.cellWidget(r, 2).isEnabled() for r in manual_rows)
    w.select_all(True)
    assert w.table.item(editable_row, 0).checkState() == Qt.Checked
    w.select_all(False)
    assert w.table.item(editable_row, 0).checkState() == Qt.Unchecked
    w.filter.setCurrentIndex(1)
    assert all(w.table.isRowHidden(r) for r, i in enumerate(w.session.result.issues) if i.severity != "high")
    assert "Căn cứ" in w.detail.text()
    w.close()


def test_cancel_and_close_worker_without_destroying_live_thread(qapp, model_dir, tmp_path):
    document = Document()
    for _ in range(80):
        document.add_paragraph("Hôm nay bạn an gì. Tôi đang di chơi.")
    path = tmp_path / "long.docx"
    document.save(path)
    w = MainWindow(model_dir)
    w.open_file(path)
    w.check()
    w.cancel()
    wait_until(qapp, lambda: w.worker is None)
    assert w.session.result is None
    assert w.notice.state == "error"
    w.close()
