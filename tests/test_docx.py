import io
import unicodedata as ud
from copy import deepcopy

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from vietspell.core.types import Issue
from vietspell.docx_io.element_check import ElementIdentifier, check_elements
from vietspell.docx_io.format_check import check_document_format, effective_font_name, inherited_font, merge_repeats
from vietspell.docx_io.reader import paragraph_leading_tab, read_docx
from vietspell.docx_io.writer import ExportMode, export_docx
from vietspell.service import DocumentSession


def saved(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def issue(text, start=0, end=None):
    end = len(text) if end is None else end
    return Issue(start, end, text[start:end], "test", "TEST")


def test_body_table_order_including_nested_table():
    document = Document()
    document.add_paragraph("Avant")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Quốc hiệu"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Bảng lồng"
    table.cell(0, 1).text = "Tiêu ngữ"
    document.add_paragraph("Nội dung")
    text = read_docx(saved(document)).text
    assert text.index("Avant") < text.index("Quốc hiệu") < text.index("Bảng lồng") < text.index("Tiêu ngữ") < text.index("Nội dung")


@pytest.mark.parametrize("mode", list(ExportMode))
def test_export_modes_cross_run_preserve_format(tmp_path, mode):
    document = Document()
    p = document.add_paragraph()
    prefix = p.add_run("Xin ")
    prefix.italic = True
    first = p.add_run("ch")
    first.bold = True
    second = p.add_run("ao")
    second.italic = True
    suffix = p.add_run(" các bạn.")
    suffix.font.size = Pt(16)
    source = saved(document)
    before = bytes(source)
    path = tmp_path / f"{mode.value}.docx"
    export_docx(source, path, [(issue("Xin chao các bạn.", 4, 8), "chào")], mode)
    result = read_docx(path)
    assert result.text == ("Xin chao các bạn." if mode == ExportMode.HIGHLIGHT else "Xin chào các bạn.")
    assert source == before
    assert result.blocks[0].runs[0].run.italic
    assert result.blocks[0].runs[-1].run.font.size.pt == 16
    root = result.document.element
    if mode == ExportMode.RED:
        assert root.xpath('.//w:color[@w:val="FF0000"]')
        assert any(s.run.bold and s.text == "chào" for s in result.blocks[0].runs)
    elif mode == ExportMode.TRACK:
        assert [n.text for n in root.iter(qn("w:delText"))] == ["ch", "ao"]
        assert len(list(root.iter(qn("w:ins")))) == 1
        assert len(list(root.iter(qn("w:del")))) == 1
        ids = [e.get(qn("w:id")) for tag in ("w:ins", "w:del") for e in root.iter(qn(tag))]
        assert len(ids) == len(set(ids))
        assert result.document.settings.element.find(qn("w:trackRevisions")) is not None
    elif mode == ExportMode.HIGHLIGHT:
        assert root.xpath('.//w:highlight[@w:val="yellow"]')
        assert any(s.text == "ch" and s.run.bold for s in result.blocks[0].runs)
        assert any(s.text == "ao" and s.run.italic for s in result.blocks[0].runs)
    for node in root.iter(qn("w:t")):
        if node.text and (node.text.startswith(" ") or node.text.endswith(" ")):
            assert node.get(qn("xml:space")) == "preserve"


def test_multiple_edits_in_one_run_and_repeated_exports(tmp_path):
    document = Document()
    document.add_paragraph("Xin  chào các bạn   .")
    source = saved(document)
    text = read_docx(source).text
    trailing = text.index("   .")
    edits = [(issue(text, 3, 5), " "), (issue(text, trailing, trailing+3), "")]
    for index in range(2):
        result = tmp_path / f"out{index}.docx"
        export_docx(source, result, edits, "PLAIN")
        assert read_docx(result).text == "Xin chào các bạn."


def test_nfd_accent_across_run_boundary_replaced_at_correct_offset(tmp_path):
    document = Document()
    p = document.add_paragraph()
    p.add_run("😀 O")
    p.add_run("\u0302ng di chơi.")
    source = saved(document)
    text = read_docx(source).text
    assert text.startswith("😀 Ông")
    start = text.index("Ông")
    out = tmp_path / "unicode.docx"
    export_docx(source, out, [(issue(text, start, start+3), "Bà")], "PLAIN")
    assert read_docx(out).text == "😀 Bà di chơi."


def test_hyperlink_text_and_suffix_are_preserved(tmp_path):
    document = Document()
    p = document.add_paragraph("Xin ")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "bookmark1")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "chao bạn"
    r.append(t)
    hyperlink.append(r)
    p._p.append(hyperlink)
    source = saved(document)
    read = read_docx(source)
    assert read.text == "Xin chao bạn"
    out = tmp_path / "link.docx"
    export_docx(source, out, [(issue(read.text, 4, 8), "chào")], "PLAIN")
    result = read_docx(out)
    assert result.text == "Xin chào bạn"
    assert len(list(result.document.element.iter(qn("w:hyperlink")))) == 1


def test_rejects_stale_and_overlapping_edits(tmp_path):
    document = Document()
    document.add_paragraph("Xin chào")
    source = saved(document)
    out = tmp_path / "safe.docx"
    out.write_bytes(b"existing output")
    with pytest.raises(ValueError):
        export_docx(source, out, [(Issue(0, 3, "Sai", "", ""), "A")])
    with pytest.raises(ValueError):
        export_docx(source, out, [(issue("Xin chào", 0, 3), "A"), (issue("Xin chào", 1, 4), "B")])
    assert out.read_bytes() == b"existing output"


def test_rejects_field_run_and_previous_revisions(tmp_path):
    for previous_revision in (False, True):
        document = Document()
        p = document.add_paragraph()
        r = p.add_run("chao")
        if previous_revision:
            ins = OxmlElement("w:ins")
            ins.set(qn("w:id"), "20")
            p._p.remove(r._r)
            ins.append(r._r)
            p._p.append(ins)
        else:
            r._r.append(OxmlElement("w:fldChar"))
        with pytest.raises(ValueError):
            export_docx(saved(document), tmp_path / "no.docx", [(issue("chao"), "chào")])


def test_tab_indentation_and_rendered_pagebreak_not_reported():
    document = Document()
    p = document.add_paragraph("\t"+"Đây là nội dung của cơ quan được soạn thảo và trình bày theo quy định. "*2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.runs[0]._r.append(OxmlElement("w:lastRenderedPageBreak"))
    read = read_docx(saved(document))
    assert paragraph_leading_tab(read.blocks[0].paragraph)
    assert not any(i.rule_id in {"FIRST_LINE_INDENT", "SOFT_BREAK"} for i in check_document_format(read, body_blocks=read.blocks))


def test_inherited_bold_size_and_theme_font():
    document = Document()
    document.styles["Normal"].font.bold = True
    document.styles["Normal"].font.size = Pt(13)
    from docx.enum.style import WD_STYLE_TYPE
    base = document.styles.add_style("ParentStyle", WD_STYLE_TYPE.PARAGRAPH)
    base.base_style = document.styles["Normal"]
    derived = document.styles.add_style("DerivedStyle", WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    p = document.add_paragraph("Thử", "DerivedStyle")
    assert inherited_font(p.runs[0], "bold", document) is True
    assert inherited_font(p.runs[0], "size", document).pt == 13
    assert effective_font_name(p.runs[0], document)  # Resolves Normal's theme, no hardcoded fallback.


def header_document():
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(13)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "BỘ CÔNG AN"
    table.cell(0, 0).add_paragraph("CÔNG AN TỈNH")
    table.cell(0, 0).add_paragraph("Số:         /CAT-ANKT")
    table.cell(0, 1).text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
    table.cell(0, 1).paragraphs[0].runs[0].font.size = Pt(12)
    table.cell(0, 1).add_paragraph("Độc lập – Tự do – Hạnh phúc")
    table.cell(0, 1).add_paragraph("Ninh Bình, ngày         tháng 10 năm 2026")
    document.add_paragraph("Kính gửi các đơn vị và cơ quan liên quan để triển khai thực hiện nhiệm vụ theo quy định của pháp luật.")
    document.add_paragraph("GIÁM ĐỐC")
    document.add_paragraph("Nguyễn Văn Nam")
    document.add_paragraph("Nơi nhận:")
    document.add_paragraph("- Như trên;")
    return document


def test_placeholder_header_elements_recognized_and_agency_not_signer():
    read = read_docx(saved(header_document()))
    findings, elements = check_elements(read)
    roles = {e.role for e in elements}
    assert {"national", "motto", "number", "date", "agency", "body", "recipients_label"} <= roles
    assert not [i for i in findings if i.rule_id.startswith("MISSING_")]
    assert not [i for i in findings if i.rule_id == "ELEMENT_TEXT_MOTTO"]
    assert not [i for i in findings if i.rule_id == "ELEMENT_SIZE_CHAIN"]
    assert not [i for i in findings if i.text == "BỘ CÔNG AN" and "BOLD" in i.rule_id]
    assert all(e.block.text != "BỘ CÔNG AN" for e in elements if e.role == "position")


def test_date_and_number_same_paragraph():
    document = header_document()
    p = document.tables[0].cell(0, 0).paragraphs[-1]
    p.add_run("\tNinh Bình, ngày 02 tháng 10 năm 2026")
    roles = [e.role for e in ElementIdentifier().identify(read_docx(saved(document)))]
    assert "number" in roles and "date" in roles


def combined_header(profile, *, date_first=False, decomposed=False, wrong_role=None, date_size=None):
    document = header_document()
    # Put both elements in one paragraph; remove the separate date to make the
    # case exercise the combined line rather than a second copy of the date.
    document.tables[0].cell(0, 1).paragraphs[-1].clear()
    paragraph = document.tables[0].cell(0, 0).paragraphs[-1]
    paragraph.clear()
    number = "Số: 123/CAT-ANKT" if profile == "HC-ND30" else "Số 05-CV/ĐU"
    date = "Ninh Bình, ngày 02 tháng 10 năm 2026"
    values = [("number", number), ("date", date)]
    if date_first:
        values.reverse()
    paragraph.add_run("  ")
    for index, (role, value) in enumerate(values):
        if index:
            paragraph.add_run("\t")
        run = paragraph.add_run(ud.normalize("NFD", value) if decomposed else value)
        run.italic = (role == "date") != (role == wrong_role)
        run.font.size = Pt(date_size if role == "date" and date_size else 13 if profile == "HC-ND30" else 14)
    return document, {"number": number, "date": date}


@pytest.mark.parametrize("profile", ["HC-ND30", "DANG-HD05"])
@pytest.mark.parametrize("date_first", [False, True])
@pytest.mark.parametrize("decomposed", [False, True])
def test_combined_header_formats_and_spans_are_independent(profile, date_first, decomposed):
    document, expected = combined_header(profile, date_first=date_first, decomposed=decomposed)
    read = read_docx(saved(document))
    findings, elements = check_elements(read, profile)
    actual = {e.role: e for e in elements if e.role in expected}
    assert set(actual) == set(expected)
    for role, value in expected.items():
        element = actual[role]
        assert element.text == value
        assert read.text[element.block.start+element.text_start:element.block.start+element.text_end] == value
    assert not [i for i in findings if i.rule_id.endswith(("_NUMBER", "_DATE")) or i.rule_id == "ELEMENT_SIZE_CHAIN"]


@pytest.mark.parametrize("profile", ["HC-ND30", "DANG-HD05"])
@pytest.mark.parametrize("wrong_role", ["number", "date"])
def test_combined_header_reports_wrong_font_on_only_the_affected_element(profile, wrong_role):
    document, expected = combined_header(profile, wrong_role=wrong_role)
    read = read_docx(saved(document))
    findings, _ = check_elements(read, profile)
    italic = [i for i in findings if i.rule_id in {"ELEMENT_ITALIC_NUMBER", "ELEMENT_ITALIC_DATE"}]
    assert [(i.rule_id, i.text) for i in italic] == [("ELEMENT_ITALIC_"+wrong_role.upper(), expected[wrong_role])]
    assert read.text[italic[0].start:italic[0].end] == expected[wrong_role]


@pytest.mark.parametrize("date_size,chain_error", [(13, False), (14, True)])
def test_combined_header_date_part_still_checks_size_chain(date_size, chain_error):
    document, _ = combined_header("HC-ND30", date_size=date_size)
    findings, _ = check_elements(read_docx(saved(document)))
    assert any(i.rule_id == "ELEMENT_SIZE_CHAIN" for i in findings) is chain_error


def test_tab_inside_date_does_not_hide_wrong_locality_font():
    document = header_document()
    paragraph = document.tables[0].cell(0, 1).paragraphs[-1]
    paragraph.clear()
    paragraph.add_run("Ninh Bình,").italic = False
    paragraph.add_run("\tngày 02 tháng 10 năm 2026").italic = True
    findings, elements = check_elements(read_docx(saved(document)))
    assert next(e.text for e in elements if e.role == "date") == paragraph.text
    assert [i.text for i in findings if i.rule_id == "ELEMENT_ITALIC_DATE"] == [paragraph.text]


def test_repeat_merging():
    result = merge_repeats([Issue(i*10, i*10+5, "mẫu", "Cùng lỗi", "X") for i in range(4)])
    assert len(result) == 1
    assert result[0].occurrences == 4 and "4 chỗ" in result[0].message


def test_session_export_no_original_overwrite_and_report(tmp_path, model_dir):
    document = Document()
    document.add_paragraph("Hôm nay bạn an gì. Tôi đang di chơi.")
    source = tmp_path / "nguồn.docx"
    source.write_bytes(saved(document))
    original = source.read_bytes()
    session = DocumentSession(source, model_dir)
    session.check()
    assert any("MẪU" in w for w in session.result.warnings)
    with pytest.raises(ValueError):
        session.export(source)
    out = tmp_path / "đã sửa.docx"
    session.export(out, mode="PLAIN")
    assert read_docx(out).text == "Hôm nay bạn ăn gì? Tôi đang đi chơi."
    assert source.read_bytes() == original
    session.report(tmp_path / "report.json")
    import json
    report = json.loads((tmp_path / "report.json").read_text(encoding="utf-8"))
    assert report["stats"]["total"] == len(report["issues"])


def test_export_reopens_the_original_snapshot(tmp_path, model_dir):
    document = Document()
    document.add_paragraph("Tôi đang di chơi.")
    source = tmp_path / "source.docx"
    source.write_bytes(saved(document))
    session = DocumentSession(source, model_dir)
    session.check()
    # External edits do not mix a new disk file with old character offsets.
    other = Document()
    other.add_paragraph("Một văn bản khác.")
    other.save(source)
    out = tmp_path / "snapshot.docx"
    session.export(out, mode="PLAIN")
    assert read_docx(out).text == "Tôi đang đi chơi."


def test_highlight_includes_manual_format_findings_without_changing_text(tmp_path, model_dir):
    document = header_document()
    source = tmp_path / "source.docx"
    document.save(source)
    session = DocumentSession(source, model_dir)
    session.check()
    assert any(not i.editable and i.end > i.start for i in session.result.issues)
    out = tmp_path / "highlight.docx"
    session.export(out, mode="HIGHLIGHT")
    highlighted = read_docx(out)
    assert highlighted.text == session.read.text
    assert highlighted.document.element.xpath('.//w:highlight[@w:val="yellow"]')


def test_letter_subject_is_12_13_regular_not_bold():
    document = header_document()
    subject = document.tables[0].cell(0, 0).add_paragraph("V/v thực hiện kế hoạch công tác")
    subject.runs[0].font.size = Pt(12)
    findings, elements = check_elements(read_docx(saved(document)))
    assert any(e.role == "letter_subject" for e in elements)
    assert not any(i.rule_id in {"ELEMENT_BOLD_LETTER_SUBJECT", "ELEMENT_SIZE_LETTER_SUBJECT"} for i in findings)


def test_national_motto_accepts_both_conventional_tone_placements():
    document = header_document()
    p = document.tables[0].cell(0, 1).paragraphs[0]
    p.text = "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM"
    findings, _ = check_elements(read_docx(saved(document)))
    assert not any(i.rule_id == "ELEMENT_TEXT_NATIONAL" for i in findings)
