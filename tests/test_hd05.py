"""Regressions grounded in the user-supplied HD05 PDF (pp. 9–13, 28, 31–33)."""
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

from vietspell.core.format_rules import FormatChecker
from vietspell.docx_io.element_check import check_elements
from vietspell.docx_io.format_check import check_document_format
from vietspell.docx_io.reader import read_docx
from vietspell.service import DocumentSession

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
from make_party_sample import make_party_sample


@pytest.mark.parametrize("letter", [False, True])
def test_hd05_valid_layout_has_no_element_or_document_findings(tmp_path, letter):
    source = tmp_path / "Văn bản Đảng.docx"
    make_party_sample(source, letter=letter)
    read = read_docx(source)
    findings, elements = check_elements(read, "DANG-HD05")
    assert not findings
    roles = {e.role for e in elements}
    assert {"party", "number", "agency", "date", "body", "authority", "position", "signer", "recipients"} <= roles
    assert not any(e.block.text.strip() == "*" for e in elements)
    assert not check_document_format(read, "DANG-HD05", [e.block for e in elements if e.role == "body"])
    if letter:
        assert {"letter_subject", "greeting", "addressees"} <= roles
        assert sum(e.role == "letter_subject" for e in elements) == 2


@pytest.mark.parametrize("letter", [False, True])
def test_hd05_detects_actual_mismatches(tmp_path, letter):
    source = tmp_path / "Sai thể thức.docx"
    make_party_sample(source, letter=letter, errors=True)
    read = read_docx(source)
    findings, elements = check_elements(read, "DANG-HD05")
    rules = {i.rule_id for i in findings}
    assert {"ELEMENT_SIZE_DATE", "ELEMENT_BOLD_POSITION", "ELEMENT_SIZE_RECIPIENTS"} <= rules
    assert "MISSING_NUMBER" not in rules
    document_issues = check_document_format(read, "DANG-HD05", [e.block for e in elements if e.role == "body"])
    assert any(i.rule_id == "FIRST_LINE_INDENT" for i in document_issues)
    if letter:
        assert {"ELEMENT_SIZE_LETTER_SUBJECT", "ELEMENT_ITALIC_LETTER_SUBJECT"} <= rules


@pytest.mark.parametrize("number,expected", [("Số 5-HD/VPTW", "Số 05-HD/VPTW"), ("Số: 05-HD/VPTW", "Số 05-HD/VPTW"), ("Số 05-HD/VPTW", None), ("Số 138-CV/BTGDVTU", None)])
def test_hd05_document_number_format(number, expected):
    issues = [i for i in FormatChecker("DANG-HD05").check(number) if i.rule_id == "PARTY_NUMBER_FORMAT"]
    assert [i.suggestions[0].text for i in issues] == ([] if expected is None else [expected])


def test_hd05_uses_its_own_date_basis_and_does_not_invent_citation_rule():
    issues = FormatChecker("DANG-HD05").check("Căn cứ điều 5 ngày 2 tháng 1 năm 2026.")
    dates = [i for i in issues if i.rule_id == "DATE_FORMAT"]
    assert len(dates) == 2
    assert all("I.4.2.2" in i.basis and "Nghị định 30" not in i.basis for i in dates)
    assert not any(i.rule_id == "CITATION_LOWERCASE" for i in issues)


def test_hd05_vv_rule_only_changes_header_prefix():
    checker = FormatChecker("DANG-HD05")
    issues = checker.check("V/v thực hiện nhiệm vụ")
    vv = next(i for i in issues if i.rule_id == "PARTY_NO_VV")
    assert vv.suggestions[0].text == "về việc "
    assert "I.5.2.3" in vv.basis
    assert not any(i.rule_id == "PARTY_NO_VV" for i in checker.check("Mẫu công văn nhà nước có chữ V/v trong trích yếu."))


def test_hd05_greeting_checks_label_and_names_separately(tmp_path):
    source = tmp_path / "Công văn.docx"
    make_party_sample(source, letter=True)
    document = Document(source)
    greeting = next(p for p in document.paragraphs if p.text.startswith("Kính gửi:"))
    greeting.runs[0].italic = False
    document.save(source)
    findings, _ = check_elements(read_docx(source), "DANG-HD05")
    wrong = [i for i in findings if i.rule_id == "ELEMENT_ITALIC_GREETING"]
    assert len(wrong) == 1 and wrong[0].text == "Kính gửi:"
    assert not any(i.rule_id == "ELEMENT_ITALIC_ADDRESSEES" for i in findings)


@pytest.mark.parametrize("letter", [False, True])
def test_hd05_session_keeps_lowercase_subject_and_removes_outdated_warning(tmp_path, model_dir, letter):
    source = tmp_path / "Bản đúng.docx"
    make_party_sample(source, letter=letter)
    result = DocumentSession(source, model_dir).check("DANG-HD05")
    assert not any(i.rule_id == "LOWERCASE_SENTENCE_START" for i in result.issues)
    assert not any("chưa đối chiếu bản gốc" in warning for warning in result.warnings)


def test_hd05_accepts_spacing_before_next_paragraph(tmp_path):
    source = tmp_path / "Giãn đoạn.docx"
    make_party_sample(source)
    document = Document(source)
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(6)
    document.save(source)
    read = read_docx(source)
    _, elements = check_elements(read, "DANG-HD05")
    findings = check_document_format(read, "DANG-HD05", [e.block for e in elements if e.role == "body"])
    assert not any(i.rule_id == "SPACE_AFTER" for i in findings)
